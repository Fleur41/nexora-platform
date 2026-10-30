"""Render a report (see services.py for its shape) as CSV, Excel or PDF."""

from __future__ import annotations

import csv
import io
from datetime import date
from decimal import Decimal

from django.utils import timezone

CONTENT_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


NUMERIC = ("money", "number", "percent")


def _csv_text(value):
    """Neutralise spreadsheet formulas in user-entered text (names, notes): CSV injection."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _xlsx_cell(sheet, row, column, value):
    """Write a cell; strings are always stored as text, never as a formula."""
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


def _number(value, kind):
    """Typed cell value for spreadsheets: money and numbers as numbers, the rest as text."""
    if value is None or value == "":
        return None
    if kind in NUMERIC:
        return Decimal(str(value))
    return value


def _display(value, kind, currency, blank=None):
    if value is None or value == "":
        return blank if blank is not None else ("—" if kind != "text" else "")
    if kind == "date":
        try:
            return date.fromisoformat(str(value)).strftime("%-d %b %Y")
        except ValueError:  # a label in a date column, e.g. "Total" on the totals row
            return str(value)
    if kind == "money":
        return f"{Decimal(str(value)):,.2f}"
    if kind == "percent":
        return f"{value}%"
    return str(value)


def _heading_lines(report) -> list[str]:
    lines = [report["organization"], report["title"], report["period"]["label"]]
    if report["property"]:
        lines.append(f"Property: {report['property']}")
    lines.append(f"Amounts in {report['currency']}")
    return lines


def to_csv(report) -> bytes:
    out = io.StringIO()
    w = csv.writer(out)
    for line in _heading_lines(report):
        w.writerow([_csv_text(line)])
    w.writerow([])
    for item in report["summary"]:
        w.writerow([item["label"], item["value"] if item["value"] is not None else ""])
    for section in report["sections"]:
        w.writerow([])
        w.writerow([section["title"]])
        cols = section["columns"]
        w.writerow([c["label"] for c in cols])
        for row in section["rows"]:
            w.writerow(
                [
                    ""
                    if row.get(c["key"]) is None
                    else row.get(c["key"])
                    if c["kind"] in NUMERIC
                    else _csv_text(row.get(c["key"]))
                    for c in cols
                ]
            )
        if section.get("totals"):
            w.writerow(
                [
                    ""
                    if section["totals"].get(c["key"]) is None
                    else section["totals"].get(c["key"])
                    for c in cols
                ]
            )
    # BOM so Excel opens UTF-8 (names, dashes) correctly.
    return ("﻿" + out.getvalue()).encode("utf-8")


def to_xlsx(report) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    navy = PatternFill("solid", fgColor="0B2545")
    bold = Font(bold=True)
    formats = {"money": "#,##0.00", "number": "#,##0", "percent": '0.0"%"'}

    ws = wb.active
    ws.title = "Summary"
    for i, line in enumerate(_heading_lines(report), start=1):
        _xlsx_cell(ws, i, 1, line).font = Font(bold=i <= 2, size=14 if i == 2 else 11)
    row = len(_heading_lines(report)) + 2
    for item in report["summary"]:
        ws.cell(row=row, column=1, value=item["label"])
        cell = ws.cell(row=row, column=2, value=_number(item["value"], item["kind"]))
        cell.number_format = formats.get(item["kind"], "General")
        cell.font = bold
        row += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 20

    used = set()
    for section in report["sections"]:
        name = section["title"][:31]
        while name in used:
            name = name[:28] + str(len(used))
        used.add(name)
        sheet = wb.create_sheet(name)
        cols = section["columns"]
        for j, c in enumerate(cols, start=1):
            cell = sheet.cell(row=1, column=j, value=c["label"])
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = navy
            cell.alignment = Alignment(horizontal="right" if c["kind"] in formats else "left")
            sheet.column_dimensions[get_column_letter(j)].width = max(
                12, min(40, len(c["label"]) + 6)
            )
        rows = list(section["rows"])
        if section.get("totals"):
            rows.append(section["totals"])
        for i, r in enumerate(rows, start=2):
            is_total = section.get("totals") is not None and i == len(rows) + 1
            for j, c in enumerate(cols, start=1):
                cell = _xlsx_cell(sheet, i, j, _number(r.get(c["key"]), c["kind"]))
                cell.number_format = formats.get(c["kind"], "General")
                if is_total:
                    cell.font = bold
        sheet.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(report) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    navy = colors.HexColor("#0B2545")
    muted = colors.HexColor("#5B6475")
    line = colors.HexColor("#E3E7EE")
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "t", parent=styles["Title"], alignment=0, textColor=navy, fontSize=18, spaceAfter=2
    )
    sub = ParagraphStyle("s", parent=styles["Normal"], textColor=muted, fontSize=9)
    h2 = ParagraphStyle(
        "h2", parent=styles["Heading2"], textColor=navy, fontSize=12, spaceBefore=8, spaceAfter=3
    )
    cell = ParagraphStyle("c", parent=styles["Normal"], fontSize=8, leading=10)
    cell_right = ParagraphStyle("cr", parent=cell, alignment=2)

    wide = max((len(s["columns"]) for s in report["sections"]), default=0) > 6
    pagesize = landscape(A4) if wide else A4
    width = pagesize[0] - 28 * mm
    generated = timezone.localtime().strftime("%d %b %Y %H:%M")

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(muted)
        canvas.drawString(
            14 * mm,
            8 * mm,
            f"{report['organization']} · {report['title']} · {report['period']['label']}",
        )
        canvas.drawRightString(
            pagesize[0] - 14 * mm, 8 * mm, f"Generated {generated} · page {doc.page}"
        )
        canvas.restoreState()

    story = [
        Paragraph(report["title"], title),
        Paragraph(
            " · ".join(
                x
                for x in (
                    report["organization"],
                    report["period"]["label"],
                    f"Property: {report['property']}" if report["property"] else "All properties",
                    f"Amounts in {report['currency']}",
                )
            ),
            sub,
        ),
        Spacer(1, 6),
    ]
    if report["summary"]:
        cards = [
            [Paragraph(i["label"], sub) for i in report["summary"]],
            [
                Paragraph(
                    f"<b>{_display(i['value'], i['kind'], report['currency'])}</b>",
                    styles["Normal"],
                )
                for i in report["summary"]
            ],
        ]
        t = Table(cards, colWidths=[width / len(report["summary"])] * len(report["summary"]))
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F5F7FA")),
                    ("BOX", (0, 0), (-1, -1), 0.5, line),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story += [t, Spacer(1, 4)]

    for section in report["sections"]:
        story.append(Paragraph(section["title"], h2))
        if section.get("note"):
            story.append(Paragraph(section["note"], sub))
            story.append(Spacer(1, 3))
        cols = section["columns"]
        numeric = [c["kind"] in ("money", "number", "percent") for c in cols]
        header = [
            Paragraph(f"<font color='white'><b>{c['label']}</b></font>", cell_right if n else cell)
            for c, n in zip(cols, numeric, strict=True)
        ]
        body = []
        rows = list(section["rows"])
        if section.get("totals"):
            rows.append(section["totals"])
        totals_row = section.get("totals")
        for r in rows:
            body.append(
                [
                    Paragraph(
                        # Totals leave unrelated columns empty rather than showing a dash.
                        _display(
                            r.get(c["key"]),
                            c["kind"],
                            report["currency"],
                            "" if r is totals_row else None,
                        ),
                        cell_right if n else cell,
                    )
                    for c, n in zip(cols, numeric, strict=True)
                ]
            )
        if not section["rows"]:
            body = [[Paragraph("Nothing to show for this period.", sub)] + [""] * (len(cols) - 1)]
        t = Table([header, *body], colWidths=[width / len(cols)] * len(cols), repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), navy),
            ("LINEBELOW", (0, 1), (-1, -1), 0.25, line),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        if section.get("totals") and section["rows"]:
            style += [
                ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F5F7FA")),
                ("LINEABOVE", (0, -1), (-1, -1), 0.75, navy),
            ]
        if not section["rows"]:
            style.append(("SPAN", (0, 1), (-1, 1)))
        t.setStyle(TableStyle(style))
        story.append(t)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=pagesize,
        leftMargin=14 * mm,
        rightMargin=14 * mm,
        topMargin=14 * mm,
        bottomMargin=16 * mm,
        title=f"{report['title']} · {report['period']['label']}",
        author=report["organization"],
    )
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


RENDERERS = {"csv": to_csv, "xlsx": to_xlsx, "pdf": to_pdf}
