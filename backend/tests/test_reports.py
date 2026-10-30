import csv
import io
from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from apps.audit.models import AuditLog
from apps.billing.models import Invoice
from apps.organizations.roles import Role
from apps.properties.models import Unit
from apps.reports import services as report_services

from .factories import LeaseFactory, PropertyFactory, UnitFactory, UnitTypeFactory

pytestmark = pytest.mark.django_db

AUG = date(2026, 8, 1)
YEAR = {"start": "2026-01-01", "end": "2026-12-31"}


def ok(res):
    assert 200 <= res.status_code < 300, getattr(res, "data", res)
    return res


@pytest.fixture
def books(make_member, org_a):
    """Two properties with a lease each: two months billed, payments, a credit, a deposit, a cost.

    A (Kilimani, 30,000 rent) and B (Ruaka, 10,000 rent), both from 1 Aug 2026, plus 1,000
    security a month. August and September are billed (in October, like a late landlord).
    """
    owner, client = make_member(org_a)
    kilimani = PropertyFactory(organization=org_a, name="Kilimani")
    ruaka = PropertyFactory(organization=org_a, name="Ruaka")
    a = LeaseFactory(
        unit=UnitFactory(unit_type=UnitTypeFactory(property=kilimani), number="A1"),
        start_date=AUG,
        rent=Decimal("30000"),
        deposit_amount=Decimal("30000"),
        rent_due_day=1,
    )
    b = LeaseFactory(
        unit=UnitFactory(unit_type=UnitTypeFactory(property=ruaka), number="B1"),
        start_date=AUG,
        rent=Decimal("10000"),
        rent_due_day=1,
    )
    UnitFactory(unit_type=UnitTypeFactory(property=ruaka), number="B2")  # vacant
    from apps.billing.defaults import ensure_default_charge_types

    security = next(ct for ct in ensure_default_charge_types(org_a) if ct.code == "security")
    ok(
        client.post(
            "/api/v1/charge-rules/",
            {
                "charge_type": str(security.pk),
                "scope": "organization",
                "method": "fixed_monthly",
                "amount": "1000",
                "effective_from": "2026-01-01",
            },
        )
    )
    for period in ("2026-08", "2026-09"):
        ok(client.post("/api/v1/billing-runs/", {"period": period}))
    aug_a = Invoice.objects.get(lease=a, period_start=AUG)
    ok(client.post(f"/api/v1/invoices/{aug_a.pk}/credit/", {"amount": "500", "reason": "Goodwill"}))
    for lease, amount, when, method in (
        (a, "31000", "2026-09-05", "mpesa"),
        (b, "5000", "2026-09-06", "cash"),
    ):
        ok(
            client.post(
                "/api/v1/payments/",
                {
                    "lease": str(lease.pk),
                    "amount": amount,
                    "received_on": when,
                    "method": method,
                    "reference": f"REF{amount}" if method == "mpesa" else "",
                    "payer_name": "",
                    "payer_phone": "",
                    "notes": "",
                },
            )
        )
    ok(
        client.post(
            f"/api/v1/deposits/{a.pk}/transactions/",
            {
                "kind": "received",
                "amount": "30000",
                "transaction_date": "2026-08-01",
                "method": "cash",
                "reference": "",
                "reason": "",
            },
        )
    )
    repair = ok(
        client.post(
            "/api/v1/maintenance/",
            {
                "unit": str(a.unit_id),
                "category": "plumbing",
                "title": "Leak",
                "description": "Sink",
                "urgency": "normal",
            },
        )
    ).data
    ok(
        client.post(
            f"/api/v1/maintenance/{repair['id']}/costs/",
            {
                "kind": "expense",
                "description": "New trap",
                "amount": "2500",
                "incurred_on": "2026-09-15",
            },
        )
    )
    return {"owner": owner, "client": client, "a": a, "b": b, "kilimani": kilimani, "ruaka": ruaka}


def report(client, key, **params):
    return ok(client.get(f"/api/v1/reports/{key}/", params)).data


def summary(data):
    return {i["label"]: i["value"] for i in data["summary"]}


def test_financial_summary_counts_income_by_the_period_it_bills_for(books):
    data = report(books["client"], "financial-summary", **YEAR)
    s = summary(data)
    # 2 months × (30,000 + 10,000 rent + 2 × 1,000 security) = 84,000, less a 500 credit.
    assert s["Net billed"] == "83500.00"
    assert s["Received"] == "36000.00"
    assert s["Collection rate"] == "43.1"
    assert s["Owed at period end"] == "47500.00"
    assert s["Maintenance costs"] == "2500.00"
    assert s["Deposits held at period end"] == "30000.00"
    income = {r["category"]: r for r in data["sections"][0]["rows"]}
    # The 500 credit on a 30,000 rent + 1,000 security invoice: 483.87 off rent, 16.13 off security.
    assert income["Rental income"]["net"] == "79516.13"
    assert income["Service charges"]["net"] == "3983.87"
    assert data["sections"][0]["totals"]["net"] == "83500.00"
    received = {r["method"]: r["amount"] for r in data["sections"][1]["rows"]}
    assert received == {"M-Pesa": "31000.00", "Cash": "5000.00"}

    # Invoiced in October, but September's rent is September's income.
    sept = summary(
        report(books["client"], "financial-summary", start="2026-09-01", end="2026-09-30")
    )
    assert sept["Net billed"] == "42000.00"
    oct_ = summary(
        report(books["client"], "financial-summary", start="2026-10-01", end="2026-10-31")
    )
    assert oct_["Net billed"] == "-500.00"  # only the credit, given in October

    chart = {p["month"]: p for p in report(books["client"], "financial-summary", **YEAR)["chart"]}
    assert chart["2026-08"]["billed"] == "42000.00" and chart["2026-09"]["collected"] == "36000.00"


def test_rent_roll(books):
    data = report(books["client"], "rent-roll", **YEAR)
    rows = {r["unit"]: r for r in data["sections"][0]["rows"]}
    assert (
        rows["A1"]["billed"],
        rows["A1"]["paid"],
        rows["A1"]["balance"],
        rows["A1"]["deposit"],
    ) == (
        "61500.00",
        "31000.00",
        "30500.00",
        "30000.00",
    )
    assert rows["B1"]["balance"] == "17000.00"
    assert data["sections"][0]["totals"]["balance"] == "47500.00"


def test_arrears_are_aged_by_due_date(books, monkeypatch):
    monkeypatch.setattr(report_services, "organization_today", lambda org: date(2026, 10, 30))
    data = report(books["client"], "arrears")
    rows = {r["unit"]: r for r in data["sections"][0]["rows"]}
    # A: credit and payment cleared August; September (due 1 Sep, 59 days late) is left.
    assert rows["A1"]["d60"] == "30500.00" and rows["A1"]["d90"] == "0.00"
    # B: 5,000 paid towards August (due 1 Aug, 90 days late); September untouched.
    assert rows["B1"]["d90"] == "6000.00" and rows["B1"]["d60"] == "11000.00"
    s = summary(data)
    assert s["Total owed"] == "47500.00" and s["Tenancies in arrears"] == "2"


def test_occupancy_and_vacant_units(books):
    Unit.objects.filter(number="B2").update(status=Unit.Status.AVAILABLE)
    data = report(books["client"], "occupancy", start="2026-08-01", end="2026-08-31")
    s = summary(data)
    assert (s["Units"], s["Occupied now"], s["Vacant now"]) == ("3", "2", "1")
    assert s["Occupancy now"] == "66.7"
    by_property = {r["property"]: r for r in data["sections"][0]["rows"]}
    assert by_property["Ruaka"]["occupancy_period"] == "50.0"  # one of two units let all August
    assert [r["unit"] for r in data["sections"][1]["rows"]] == ["B2"]


def test_deposits_and_maintenance_costs(books):
    deposits = report(books["client"], "deposits", **YEAR)
    assert summary(deposits)["Held for tenants"] == "30000.00"
    costs = report(books["client"], "maintenance-costs", **YEAR)
    assert summary(costs)["Spent"] == "2500.00"
    assert costs["sections"][0]["rows"][0]["supplier"] == ""


def test_property_filter_and_assigned_scope(books, make_member, org_a):
    only_ruaka = summary(
        report(books["client"], "financial-summary", property=str(books["ruaka"].pk), **YEAR)
    )
    assert only_ruaka["Net billed"] == "22000.00" and only_ruaka["Received"] == "5000.00"

    manager, client = make_member(org_a, role=Role.PROPERTY_MANAGER)
    ok(
        books["client"].post(
            "/api/v1/property-assignments/",
            {"property": str(books["kilimani"].pk), "membership": str(manager.pk)},
        )
    )
    mine = summary(report(client, "financial-summary", **YEAR))
    assert mine["Net billed"] == "61500.00"  # Kilimani only
    res = client.get("/api/v1/reports/financial-summary/", {"property": str(books["ruaka"].pk)})
    assert res.status_code == 404


def test_permissions(books, make_member, org_a, org_b):
    _, caretaker = make_member(org_a, role=Role.CARETAKER)
    assert caretaker.get("/api/v1/reports/occupancy/").status_code == 403
    _, accountant = make_member(org_a, role=Role.ACCOUNTANT)
    ok(accountant.get("/api/v1/reports/rent-roll/"))
    _, outsider = make_member(org_b)
    assert summary(report(outsider, "financial-summary", **YEAR))["Net billed"] == "0.00"
    assert books["client"].get("/api/v1/reports/nope/").status_code == 404
    res = books["client"].get(
        "/api/v1/reports/rent-roll/", {"start": "2026-09-01", "end": "2026-08-01"}
    )
    assert res.status_code == 400


def test_exports_csv_xlsx_pdf_and_audit(books):
    client = books["client"]
    res = client.get("/api/v1/reports/rent-roll/", {**YEAR, "export": "csv"})
    assert res.status_code == 200 and res["Content-Type"].startswith("text/csv")
    assert 'filename="rent-roll-2026-01-01-to-2026-12-31.csv"' in res["Content-Disposition"]
    rows = list(csv.reader(io.StringIO(res.content.decode("utf-8-sig"))))
    assert rows[[r[:1] for r in rows].index(["Kilimani"])][:2] == ["Kilimani", "A1"]

    res = client.get("/api/v1/reports/financial-summary/", {**YEAR, "export": "xlsx"})
    wb = load_workbook(io.BytesIO(res.content))
    assert wb.sheetnames[:2] == ["Summary", "Income billed"]
    assert wb["Income billed"]["A1"].value == "Category"

    res = client.get("/api/v1/reports/arrears/", {"export": "pdf"})
    assert res["Content-Type"] == "application/pdf" and res.content.startswith(b"%PDF")

    exports = AuditLog.objects.filter(action="report.exported")
    assert sorted(e.metadata["format"] for e in exports) == ["csv", "pdf", "xlsx"]


def test_report_list_marks_cost_report_unavailable_without_cost_access(make_member, org_a):
    membership, client = make_member(org_a, role=Role.ACCOUNTANT)
    keys = {r["key"]: r["available"] for r in ok(client.get("/api/v1/reports/")).data}
    assert keys["maintenance-costs"] is True  # accountants see costs by default
    membership.revoked_permissions = ["maintenance.costs"]
    membership.save()
    res = client.get("/api/v1/reports/maintenance-costs/")
    assert res.status_code == 403
    assert "Maintenance costs" not in summary(report(client, "financial-summary"))


def test_exports_never_turn_user_text_into_formulas():
    import io

    from openpyxl import load_workbook

    from apps.reports import exporters

    report = {
        "organization": "=cmd|'/C calc'!A0",
        "title": "Rent roll",
        "period": {"label": "October 2026"},
        "property": None,
        "currency": "KES",
        "summary": [],
        "sections": [
            {
                "title": "Units",
                "columns": [
                    {"key": "tenant", "label": "Tenant", "kind": "text"},
                    {"key": "balance", "label": "Balance", "kind": "money"},
                ],
                "rows": [{"tenant": '=HYPERLINK("http://evil.example","x")', "balance": "-500"}],
            }
        ],
    }
    csv_text = exporters.to_csv(report).decode("utf-8-sig")
    assert "'=HYPERLINK" in csv_text and "'=cmd" in csv_text
    assert ",-500" in csv_text  # numbers keep their sign
    wb = load_workbook(io.BytesIO(exporters.to_xlsx(report)))
    assert wb["Summary"]["A1"].data_type == "s"
    cell = wb["Units"]["A2"]
    assert cell.data_type == "s" and cell.value.startswith("=HYPERLINK")
    assert wb["Units"]["B2"].value == -500
